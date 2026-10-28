"""An Excel report of the system's labels against the practice labels.

    python -m hype.report                       # -> report.xlsx
    python -m hype.report --out my_report.xlsx
    python -m hype.report --table results.csv   # also list the rows of a table that are not practice rows

Two labels for each practice row, because the number that matters depends on what
the system saw:

- **with all knowledge**: the row labelled with every other practice row known,
  as in `results.csv`. The resolver learned from these rows too, so this is an
  upper bound.
- **on unseen pages**: the row's page left out of everything the system learned
  and knows, five folds of pages, as the arena run sees new rows. This is the
  honest number.

Rows where the system differs from the practice label are shaded; every label
cell is coloured by its label. Needs `openpyxl` (pip install openpyxl), which the
system itself does not.
"""

from __future__ import annotations

import argparse
import contextlib
import io
import sys
from collections import Counter

from .characterize import DEFAULT_LEXICON, DEFAULT_PRACTICE
from .label import LABELS, Labeller, out_of_fold_evidence
from .lexicon import Lexicon
from .pipeline import HypePipeline
from .resolve import read_rows, segment_pages
from .run import as_json, label_table, read_table

COLOURS = {"true": "C6EFCE", "false": "FFC7CE", "unsure": "FFEB9C"}
MISMATCH = "F8D7DA"
HEADER = "1F3B57"


def unseen_labels(lexicon: Lexicon, rows: list[dict], folds: int = 5) -> dict[str, tuple[str, float | None]]:
    """Each practice row labelled by a system that never saw its page: learned and known without it."""
    with contextlib.redirect_stderr(io.StringIO()):
        evidence, labels, fold = out_of_fold_evidence(lexicon, segment_pages(rows), folds=folds)
    out = {}
    for k in range(folds):
        labeller = Labeller()
        labeller.fit([e for e, j in zip(evidence, fold) if j != k], [y for y, j in zip(labels, fold) if j != k])
        test = [i for i, j in enumerate(fold) if j == k]
        for i, (label, conf, _) in zip(test, labeller.predict_all([evidence[i] for i in test])):
            out[evidence[i].id] = (label, conf)
    return out


def scores(pairs: list[tuple[str, str]]) -> dict:
    """Accuracy, per-label precision and recall, and the confusion counts of (truth, predicted) pairs."""
    confusion = Counter(pairs)
    right = sum(n for (t, p), n in confusion.items() if t == p)
    per = {}
    for label in LABELS:
        tp = confusion[(label, label)]
        predicted = sum(n for (t, p), n in confusion.items() if p == label)
        actual = sum(n for (t, p), n in confusion.items() if t == label)
        per[label] = (tp / predicted if predicted else None, tp / actual if actual else None, actual)
    return {"accuracy": right / len(pairs), "right": right, "rows": len(pairs), "per": per, "confusion": confusion}


def build(out_path: str, table: str | None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    practice = read_rows(DEFAULT_PRACTICE)
    truth = {r["id"]: r["label"] for r in practice}
    print("labelling with all knowledge...", file=sys.stderr, flush=True)
    pipeline = HypePipeline.load(DEFAULT_LEXICON)
    rows, verdicts, readings = label_table(pipeline, [dict(r) for r in practice], relabel=True)
    print("labelling each page unseen (five folds, about a minute and a half)...", file=sys.stderr, flush=True)
    unseen = unseen_labels(Lexicon.load(DEFAULT_LEXICON), practice)

    extra = []
    if table:
        extra = [r for r in read_table(table) if r["id"] not in truth]

    full = scores([(truth[r["id"]], verdicts[r["id"]].label) for r in rows])
    honest = scores([(truth[i], unseen[i][0]) for i in truth if i in unseen])

    wb = Workbook()
    fill = lambda hex_: PatternFill("solid", start_color=hex_, end_color=hex_)
    bold = Font(bold=True)
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D0D7DE")

    # --- summary ---------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Hype Arena: the system's labels against the practice labels"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "", "with all knowledge", "on unseen pages (honest)", "always 'true' (baseline)"
    for c in "ABCD":
        ws[f"{c}3"].font = head_font
        ws[f"{c}3"].fill = fill(HEADER)
    baseline = sum(1 for v in truth.values() if v == "true") / len(truth)
    ws.append(["accuracy", full["accuracy"], honest["accuracy"], baseline])
    ws.append(["rows right", f"{full['right']} of {full['rows']}", f"{honest['right']} of {honest['rows']}",
               f"{sum(1 for v in truth.values() if v == 'true')} of {len(truth)}"])
    for label in LABELS:
        for name, k in (("precision", 0), ("recall", 1)):
            ws.append([f"{label}: {name}", full["per"][label][k], honest["per"][label][k], None])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=2, max_col=4):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0%"
    r0 = ws.max_row + 2
    ws.cell(r0, 1, "Confusion on unseen pages: rows are the practice label, columns the system's").font = bold
    ws.cell(r0 + 1, 1, "practice \\ system").font = head_font
    ws.cell(r0 + 1, 1).fill = fill(HEADER)
    for j, label in enumerate(LABELS, start=2):
        ws.cell(r0 + 1, j, label).font = head_font
        ws.cell(r0 + 1, j).fill = fill(HEADER)
    for i, t in enumerate(LABELS, start=r0 + 2):
        ws.cell(i, 1, t).fill = fill(COLOURS[t])
        for j, p in enumerate(LABELS, start=2):
            c = ws.cell(i, j, honest["confusion"][(t, p)])
            if t == p:
                c.font = bold
    note = ws.max_row + 2
    lines = [
        "With all knowledge: each row is labelled with every other practice row known, and the resolver learned from them; an upper bound.",
        "On unseen pages: each page is left out of what the system learned and knows, as new rows in the arena run. This is the number to trust.",
        "Sheet 'Rows': a shaded row is one where the system differs from the practice label.",
    ]
    for k, text in enumerate(lines):
        ws.cell(note + k, 1, text)
    ws.column_dimensions["A"].width = 34
    for c in "BCD":
        ws.column_dimensions[c].width = 26

    # --- rows --------------------------------------------------------------------------------
    rs = wb.create_sheet("Rows")
    columns = ["id", "date", "text", "practice label", "system label", "confidence", "same?",
               "label on unseen pages", "same on unseen?", "decided by", "form", "residue"]
    rs.append(columns)
    for j in range(1, len(columns) + 1):
        c = rs.cell(1, j)
        c.font = head_font
        c.fill = fill(HEADER)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        v = verdicts[r["id"]]
        u_label, u_conf = unseen.get(r["id"], ("", None))
        t = truth[r["id"]]
        rs.append([r["id"], r["date"], r["text"], t, v.label, v.confidence, "yes" if v.label == t else "NO",
                   u_label, "yes" if u_label == t else "NO", v.decided_by,
                   as_json(readings[r["id"]].form), as_json(readings[r["id"]].residue)])
        n = rs.max_row
        if v.label != t:
            for j in range(1, len(columns) + 1):
                rs.cell(n, j).fill = fill(MISMATCH)
        for j, label in ((4, t), (5, v.label), (8, u_label)):
            if label in COLOURS:
                rs.cell(n, j).fill = fill(COLOURS[label])
        for j in (7, 9):
            if rs.cell(n, j).value == "NO":
                rs.cell(n, j).font = Font(bold=True, color="9C0006")
    # the other rows of a table, such as the sentences added with --add: no practice label to compare
    if extra:
        verdicts_extra, readings_extra = pipeline.analyse([{**r, "label": ""} for r in extra], same_page=False)
        by_id = {x.id: x for x in verdicts_extra}
        for r in extra:
            v = by_id.get(r["id"])
            rs.append([r["id"], r["date"], r["text"], "", v.label if v else "", v.confidence if v else None, "",
                       "", "", v.decided_by if v else "", as_json(readings_extra[r["id"]].form),
                       as_json(readings_extra[r["id"]].residue)])
            if v and v.label in COLOURS:
                rs.cell(rs.max_row, 5).fill = fill(COLOURS[v.label])
    widths = [9, 11, 70, 13, 13, 11, 8, 13, 10, 18, 60, 40]
    for j, w in enumerate(widths, start=1):
        rs.column_dimensions[get_column_letter(j)].width = w
    for row in rs.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (3, 11, 12))
            cell.border = Border(bottom=thin)
        row[5].number_format = "0.00"
    rs.freeze_panes = "D2"
    rs.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{rs.max_row}"

    wb.save(out_path)
    print(f"accuracy with all knowledge: {full['accuracy']:.1%} ({full['right']} of {full['rows']})")
    print(f"accuracy on unseen pages:    {honest['accuracy']:.1%} ({honest['right']} of {honest['rows']})")
    print(f"always 'true' baseline:      {baseline:.1%}")
    print(f"wrote {out_path}")


def results_to_excel(csv_path: str, out_path: str) -> None:
    """A table the system wrote (`results.csv`), as Excel: its own columns, the practice label beside
    the system's where the row is a practice row, label cells coloured, differing rows shaded."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    truth = {r["id"]: r["label"] for r in read_rows(DEFAULT_PRACTICE)}
    rows = read_table(csv_path)
    columns = list(rows[0].keys()) if rows else []
    at = columns.index("label") + 1 if "label" in columns else len(columns)
    columns = columns[:at] + ["practice label", "same?"] + columns[at:]
    fill = lambda hex_: PatternFill("solid", start_color=hex_, end_color=hex_)
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(columns)
    for j in range(1, len(columns) + 1):
        ws.cell(1, j).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, j).fill = fill(HEADER)
    compared = right = 0
    for r in rows:
        t = truth.get(r["id"], "")
        same = "" if not t else ("yes" if r.get("label") == t else "NO")
        if t:
            compared += 1
            right += same == "yes"
        ws.append([{"practice label": t, "same?": same}.get(c, r.get(c, "")) for c in columns])
        n = ws.max_row
        if same == "NO":
            for j in range(1, len(columns) + 1):
                ws.cell(n, j).fill = fill(MISMATCH)
            ws.cell(n, columns.index("same?") + 1).font = Font(bold=True, color="9C0006")
        for name in ("label", "practice label"):
            if name in columns and ws.cell(n, columns.index(name) + 1).value in COLOURS:
                ws.cell(n, columns.index(name) + 1).fill = fill(COLOURS[ws.cell(n, columns.index(name) + 1).value])
    widths = {"id": 9, "text": 70, "date": 11, "label": 9, "practice label": 13, "same?": 7, "form": 60, "residue": 40}
    for j, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 16)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=columns[cell.column - 1] in ("text", "form", "residue"))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{ws.max_row}"
    wb.save(out_path)
    if compared:
        print(f"{right} of {compared} practice rows have the practice label ({right / compared:.1%}); "
              f"{compared - right} rows differ and are shaded")
    print(f"wrote {out_path} ({len(rows)} rows)")


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="An Excel report of the system's labels against the practice labels.")
    parser.add_argument("--out", help="the workbook to write (default: report.xlsx, or results.xlsx with --results)")
    parser.add_argument("--table", help="a table (such as results.csv) whose non-practice rows are listed too")
    parser.add_argument("--results", help="only turn this table (such as results.csv) into Excel, with differences shaded")
    args = parser.parse_args(argv)
    for stream in (sys.stdout, sys.stderr):
        with contextlib.suppress(AttributeError, ValueError):
            stream.reconfigure(encoding="utf-8", errors="replace")
    if args.results:
        results_to_excel(args.results, args.out or args.results.rsplit(".", 1)[0] + ".xlsx")
    else:
        build(args.out or "report.xlsx", args.table)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
